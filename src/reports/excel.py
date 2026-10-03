"""
Excel workbook builder shared by reports and the Excel backup.

Each sheet gets: a title block, a bold header row with filters and frozen
panes, Indian-style rupee number format on money columns, and a TOTAL row
built from SUM formulas (so totals stay correct if rows are edited).
"""
from io import BytesIO
from typing import Dict, Iterable, Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from helper.clock import now_ist

FONT = "Arial"
MONEY_FORMAT = '[>=10000000]##\\,##\\,##\\,##0.00;[>=100000]##\\,##\\,##0.00;##,##0.00'
HEADER_FILL = PatternFill("solid", fgColor="DCE6F1")
_INVALID_SHEET_CHARS = str.maketrans({c: " " for c in "[]:*?/\\"})


def _safe_sheet_name(name: str, used: set) -> str:
    base = name.translate(_INVALID_SHEET_CHARS).strip()[:31] or "Sheet"
    candidate, counter = base, 2
    while candidate.lower() in used:
        suffix = f" ({counter})"
        candidate = base[: 31 - len(suffix)] + suffix
        counter += 1
    used.add(candidate.lower())
    return candidate


def _cell_value(value):
    if value is None:
        return None
    if isinstance(value, float) and pd.isna(value):
        return None
    if value is pd.NaT:
        return None
    if hasattr(value, "item"):  # numpy scalars
        return value.item()
    return value


def _write_sheet(workbook: Workbook, name: str, frame: pd.DataFrame, title: str, subtitle: str,
                 money_columns: Iterable[str], add_totals: bool, used: set) -> None:
    sheet = workbook.create_sheet(_safe_sheet_name(name, used))
    money = [column for column in frame.columns if column in set(money_columns)]

    sheet["A1"] = title
    sheet["A1"].font = Font(name=FONT, bold=True, size=13)
    sheet["A2"] = subtitle
    sheet["A2"].font = Font(name=FONT, italic=True, size=9, color="555555")
    header_row = 4

    for col_index, column in enumerate(frame.columns, start=1):
        cell = sheet.cell(row=header_row, column=col_index, value=str(column))
        cell.font = Font(name=FONT, bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for row_offset, row in enumerate(frame.itertuples(index=False), start=1):
        for col_index, value in enumerate(row, start=1):
            cell = sheet.cell(row=header_row + row_offset, column=col_index, value=_cell_value(value))
            cell.font = Font(name=FONT)
            if frame.columns[col_index - 1] in money:
                cell.number_format = MONEY_FORMAT

    last_data_row = header_row + len(frame)
    if add_totals and money and len(frame):
        total_row = last_data_row + 1
        label = sheet.cell(row=total_row, column=1, value="TOTAL")
        label.font = Font(name=FONT, bold=True)
        for column in money:
            col_index = list(frame.columns).index(column) + 1
            letter = get_column_letter(col_index)
            cell = sheet.cell(row=total_row, column=col_index,
                              value=f"=SUM({letter}{header_row + 1}:{letter}{last_data_row})")
            cell.font = Font(name=FONT, bold=True)
            cell.number_format = MONEY_FORMAT

    for col_index, column in enumerate(frame.columns, start=1):
        values = [str(column)] + [str(v) for v in frame.iloc[:200, col_index - 1].tolist()]
        sheet.column_dimensions[get_column_letter(col_index)].width = min(max(len(v) for v in values) + 3, 50)
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
    if len(frame.columns):
        sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(frame.columns))}{max(last_data_row, header_row)}"


def build_workbook(sheets: Dict[str, pd.DataFrame], title: str, school_name: str = "",
                   money_columns: Optional[Iterable[str]] = None, add_totals: bool = True) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    subtitle = f"{school_name + ' | ' if school_name else ''}Generated {now_ist():%d %b %Y %H:%M} IST"
    used: set = set()
    for name, frame in sheets.items():
        _write_sheet(workbook, name, frame, f"{title} - {name}", subtitle, money_columns or [], add_totals, used)
    if not sheets:
        workbook.create_sheet("Empty")["A1"] = "No data"
    workbook.calculation.fullCalcOnLoad = True
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
