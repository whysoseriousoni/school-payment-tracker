import sqlite3
from datetime import datetime, timedelta

import pytest

from backup_and_restore import service
from config import settings
from data_management.dto.payment import VoidRequest
from data_management.services import payment_service
from data_management.services.errors import BusinessRuleError


@pytest.fixture
def backups(app_db, tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "BACKUP_DIR", tmp_path / "app_backups")
    return tmp_path / "app_backups"


def _counts(path):
    conn = sqlite3.connect(path)
    try:
        return {t: conn.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0] for t in service.TABLE_ORDER}
    finally:
        conn.close()


def test_weekly_backup_only_when_due_and_pruned(backups, monkeypatch):
    first = service.run_scheduled_backup_if_due()
    assert first is not None and first.exists()
    assert service.run_scheduled_backup_if_due() is None
    assert service.run_scheduled_backup_if_due(now=datetime.now() + timedelta(days=8)) is not None
    monkeypatch.setattr(settings, "AUTO_BACKUP_KEEP", 1)
    assert service.prune_automatic_backups(keep=1) == 1
    assert len(service.list_backups()) == 1


def test_db_backup_restore_round_trip(app_db, backups, connect):
    before = _counts(app_db)
    snapshot = service.create_backup("manual")
    conn = connect(app_db)
    conn.execute("UPDATE student SET name = 'CHANGED' WHERE id = 1")
    conn.commit()

    result = service.restore_from_db_file(snapshot.read_bytes())
    assert result.safety_backup.exists() and "pre_restore" in result.safety_backup.name
    assert connect(app_db).execute("SELECT name FROM student WHERE id = 1").fetchone()[0] == "PRADEEP"
    assert _counts(app_db) == before


def test_db_restore_rejects_bad_files(backups):
    with pytest.raises(BusinessRuleError, match="Not a valid database"):
        service.restore_from_db_file(b"this is not sqlite" * 100)


def test_excel_backup_restore_round_trip_with_voided_receipt(app_db, backups, connect):
    payment_service.void_payment(VoidRequest(payment_id=1, reason="Duplicate entry"), "admin")
    replacement_due = connect(app_db).execute(
        "SELECT fee_due_id FROM v_fee_due_status WHERE student_id = 1 AND fee_type = 'TUITION'").fetchone()[0]
    from datetime import date

    from data_management.dto.payment import AllocationInput, PaymentCreate
    payment_service.create_payment(PaymentCreate(
        student_id=1, enrollment_id=1, paid_on=date(2026, 9, 1), payment_method="UPI", billing_name="Latha",
        allocations=[AllocationInput(fee_due_id=replacement_due, amount_paise=250000)]), "admin")
    before = _counts(app_db)
    ledger_before = connect(app_db).execute("SELECT * FROM v_fee_due_status ORDER BY fee_due_id").fetchall()
    workbook = service.export_excel_backup()

    conn = connect(app_db)
    conn.execute("DELETE FROM student_guardian")
    conn.commit()
    result = service.restore_from_excel(workbook)
    assert result.rows["payment"] == 4 and result.rows["payment_allocation"] == before["payment_allocation"]
    assert _counts(app_db) == before
    restored = connect(app_db)
    assert [tuple(r) for r in restored.execute("SELECT * FROM v_fee_due_status ORDER BY fee_due_id")] == \
        [tuple(r) for r in ledger_before]
    voided = restored.execute("SELECT is_voided, void_reason FROM payment WHERE id = 1").fetchone()
    assert tuple(voided) == (1, "Duplicate entry")
    with pytest.raises(sqlite3.IntegrityError, match="void them instead"):
        restored.execute("DELETE FROM payment WHERE id = 2")


def test_excel_restore_rejects_tampered_money(app_db, backups, connect):
    import io

    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(service.export_excel_backup()))
    sheet = workbook["payment_allocation"]
    amount_col = [c.value for c in sheet[1]].index("amount_paise") + 1
    sheet.cell(row=2, column=amount_col, value=99999999)
    buffer = io.BytesIO()
    workbook.save(buffer)
    before = _counts(app_db)
    with pytest.raises(BusinessRuleError, match="failed validation"):
        service.restore_from_excel(buffer.getvalue())
    assert _counts(app_db) == before, "nothing changed after a rejected restore"


def test_excel_restore_rejects_foreign_files(backups):
    with pytest.raises(BusinessRuleError):
        service.restore_from_excel(b"junk")
