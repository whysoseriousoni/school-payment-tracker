from datetime import date
from io import BytesIO

import pytest
from openpyxl import load_workbook

from data_management.dto.fee import FeePlanInput, MilestoneInput, MilestoneSet, OneOffDueCreate
from data_management.dto.payment import AllocationInput, PaymentCreate
from data_management.services import academic_year_service, fee_service, payment_service, report_service
from data_management.services.report_service import ReportFilter
from reports.excel import build_workbook


@pytest.fixture
def year(app_db):
    """Real prototype data + UKG plans (general Rs 10,000, DG 1 Rs 5,000) and milestones 50% Oct / 100% Mar."""
    year = academic_year_service.get_current_year()
    for code, category, amount in (("UKG-FEE-1", None, 1000000), ("UKG-DG1-1", "DG 1", 500000)):
        fee_service.save_plan(FeePlanInput(academic_year_id=year.id, code=code, fee_type="TUITION",
                                           student_class="UKG", category=category, annual_amount_paise=amount,
                                           is_default=True))
    fee_service.assign_defaults(year.id)
    fee_service.save_milestones(MilestoneSet(academic_year_id=year.id, milestones=[
        MilestoneInput(due_date=date(2026, 10, 31), cumulative_percent=50),
        MilestoneInput(due_date=date(2027, 3, 31), cumulative_percent=100)]))
    return year


def _pay(enrollment_id, student_id, amount, paid_on):
    due = fee_service.get_ledger(enrollment_id)[0]
    return payment_service.create_payment(PaymentCreate(
        student_id=student_id, enrollment_id=enrollment_id, paid_on=paid_on, payment_method="UPI",
        billing_name="Parent", allocations=[AllocationInput(fee_due_id=due.fee_due_id, amount_paise=amount)]), "t")


def test_nobody_overdue_before_first_milestone(year):
    balances = report_service.student_balances(ReportFilter(year.id, as_of=date(2026, 10, 30)))
    assert set(balances["payment_status"]) == {"Paid up", "No fees set"}
    assert balances.set_index("name").loc["VBN", "fee_plan"] == "UKG-DG1-1"
    assert balances.set_index("name").loc["PEEDARP", "fee_plan"] == "not assigned"


def test_overdue_after_milestone(year):
    _pay(6, 6, 200000, date(2026, 10, 15))  # VBN pays Rs 2,000 of Rs 5,000
    balances = report_service.student_balances(ReportFilter(year.id, as_of=date(2026, 11, 5))).set_index("name")
    assert balances.loc["VBN", "expected_paise"] == 250000 and balances.loc["VBN", "overdue_paise"] == 50000
    assert balances.loc["VBN", "overdue_details"] == "Tuition Fee (UKG-DG1-1): ₹500.00"
    assert balances.loc["PRADEEP", "overdue_paise"] == 1500000 - 1130000
    assert balances.loc["VBN", "receipts"] == 1

    summary = report_service.dashboard(ReportFilter(year.id, as_of=date(2026, 11, 5)))
    assert (summary["expected_percent"], summary["enrolled"], summary["paid_up"], summary["defaulters"],
            summary["no_fees"]) == (50, 6, 0, 4, 2)
    assert summary["next_milestone"] == (date(2027, 3, 31), 100)
    assert summary["collected_paise"] == 1130000 + 200000

    dg = report_service.defaulters(ReportFilter(year.id, as_of=date(2026, 11, 5), category="DG 1"))
    assert list(dg["Student"]) == ["VBN"] and dg.iloc[0]["Overdue (Rs)"] == 500


def test_as_of_counts_only_payments_made_by_then(year):
    _pay(6, 6, 250000, date(2026, 11, 20))  # pays the 50% three weeks late
    early = report_service.student_balances(ReportFilter(year.id, as_of=date(2026, 11, 5))).set_index("name")
    late = report_service.student_balances(ReportFilter(year.id, as_of=date(2026, 11, 25))).set_index("name")
    assert early.loc["VBN", "overdue_paise"] == 250000 and early.loc["VBN", "total_paid_paise"] == 0
    assert late.loc["VBN", "payment_status"] == "Paid up"


def test_one_off_fee_expected_immediately(year):
    fee_service.add_one_off_due(OneOffDueCreate(enrollment_id=6, fee_type="UNIFORM", amount_paise=50000))
    balances = report_service.student_balances(ReportFilter(year.id)).set_index("name")
    assert balances.loc["VBN", "overdue_paise"] >= 50000


def test_end_of_term_expects_everything(year):
    full_term = report_service.student_balances(ReportFilter(year.id, as_of=date(2027, 9, 1))).set_index("name")
    assert full_term.loc["PRADEEP", "overdue_paise"] == full_term.loc["PRADEEP", "balance_paise"] == 1870000


def test_trend_breakdowns_and_comparison(year):
    report_filter = ReportFilter(year.id, as_of=date(2026, 11, 5))
    trend = report_service.monthly_trend(report_filter).set_index("Month")
    assert list(trend.index)[:2] == ["Jun 2026", "Jul 2026"] and len(trend) == 12
    assert trend.loc["Jun 2026", "Collected in month"] == 3000 and trend.loc["Jul 2026", "Collected so far"] == 11300
    assert trend.loc["Sep 2026", "Expected by month end"] == 0
    assert trend.loc["Oct 2026", "Expected by month end"] == (30000 + 10000 * 2 + 5000) / 2
    assert trend.loc["May 2027", "Expected by month end"] == 55000
    assert report_service.class_breakdown(report_filter).iloc[0]["Class"] == "UKG-A"
    assert report_service.payment_method_split(report_filter).set_index("Method").loc["CASH", "Amount"] == 11300
    assert report_service.category_breakdown(report_filter).set_index("Category").loc["DG 1", "Defaulters"] == 1
    assert report_service.term_comparison().iloc[0]["Paid"] == 11300


def _workbook(sheets):
    return load_workbook(BytesIO(build_workbook(sheets, "Test", "My School",
                                                money_columns=report_service.MONEY_COLUMNS)))


def test_excel_reports_build(year):
    _pay(6, 6, 200000, date(2026, 11, 3))
    monthly = report_service.monthly_report(date(2026, 11, 1), include_all_students=True)
    students = monthly["Students at 30 Nov 2026"].set_index("Student")
    assert students.loc["VBN", "Paid in month (Rs)"] == 2000 and students.loc["VBN", "Overdue (Rs)"] == 500
    workbook = _workbook(monthly)
    payments_sheet = workbook["Payments"]
    assert payments_sheet["A4"].value == "Receipt No"
    assert payments_sheet.cell(payments_sheet.max_row, 1).value == "TOTAL"

    july = report_service.monthly_report(date(2026, 7, 1), include_all_students=False)
    assert july["Payments"]["Amount (Rs)"].sum() == 8300 and "Students at 31 Jul 2026" not in july

    fy = report_service.financial_year_report(2026)
    assert fy["Monthly summary"]["Collected (Rs)"].sum() == 13300

    operational = report_service.operational_year_report(ReportFilter(year.id, as_of=date(2026, 11, 30)))
    grid = operational["Paid per month"].set_index("Student")
    assert grid.loc["PRADEEP", "Jul 2026 (Rs)"] == 8300 and grid.loc["PRADEEP", "Balance (Rs)"] == 18700
    summary = dict(zip(operational["Summary"]["Measure"], operational["Summary"]["Value"]))
    assert summary["Expected share of annual fees by now"] == "50%"
    _workbook(operational)

    custom = report_service.custom_report(date(2026, 6, 1), date(2026, 7, 31), methods=["CASH"], fee_types=["TUITION"])
    assert set(custom["Fee details"]["Fee"]) == {"Tuition Fee (2-FEE-1)"}
    assert custom["Fee details"]["Amount (Rs)"].sum() == 11300
    _workbook(custom)


def test_empty_workbook_is_valid():
    assert load_workbook(BytesIO(build_workbook({}, "Nothing"))).sheetnames == ["Empty"]
